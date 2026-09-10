"""Write metrics × methods Excel with the best cell per metric bold+underlined."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

HIGHER_BETTER = {
    "val_median_pcc",
    "test_median_pcc",
    "test_mean_pcc",
    "test_median_spearman",
    "test_protein_median_pcc",
    "test_metabolite_median_pcc",
}
LOWER_BETTER = {
    "test_rmse",
    "test_mae",
    "test_protein_rmse",
    "test_metabolite_rmse",
}

METRIC_ROWS = [
    "val_median_pcc",
    "test_median_pcc",
    "test_mean_pcc",
    "test_median_spearman",
    "test_rmse",
    "test_mae",
    "test_protein_median_pcc",
    "test_metabolite_median_pcc",
    "test_protein_rmse",
    "test_metabolite_rmse",
]


def _is_better(metric: str, value: float, best: float) -> bool:
    if metric in LOWER_BETTER:
        return value < best
    return value > best


def metrics_methods_frame(rows: list[dict[str, Any]]) -> pd.DataFrame:
    ok = [row for row in rows if row.get("model") and not row.get("error")]
    methods = [str(row["model"]) for row in ok]
    data = {"metric": [name for name in METRIC_ROWS]}
    for row in ok:
        data[str(row["model"])] = [row.get(name) for name in METRIC_ROWS]
    frame = pd.DataFrame(data)
    if methods:
        frame = frame.loc[:, ["metric", *methods]]
    return frame


def write_metrics_excel(rows: list[dict[str, Any]], dest: Path) -> Path:
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    frame = metrics_methods_frame(rows)
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "metrics_x_methods"
    for excel_row in dataframe_to_rows(frame, index=False, header=True):
        sheet.append(excel_row)
    methods = [col for col in frame.columns if col != "metric"]
    highlight = Font(bold=True, underline="single")
    for r_idx, metric in enumerate(METRIC_ROWS, start=2):
        best_col = None
        best_val: float | None = None
        for c_idx, method in enumerate(methods, start=2):
            raw = sheet.cell(r_idx, c_idx).value
            try:
                value = float(raw)
            except (TypeError, ValueError):
                continue
            if not np.isfinite(value):
                continue
            if best_val is None or _is_better(metric, value, best_val):
                best_val = value
                best_col = c_idx
        if best_col is not None:
            sheet.cell(r_idx, best_col).font = highlight
    book.save(dest)
    return dest
